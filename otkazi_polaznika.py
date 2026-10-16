"""
Izbaci polaznika iz tablice tečaja - kad netko otkaže dolazak.

Prije svake izmjene nastane kopija cijele tablice u podmapi 'Otkazani', s
datumom i vremenom u nazivu. Tko je izbačen, kada i koliko je do tada
uplatio, zapisuje se u 'otkazani.json'. Dvije mreže ispod iste trapeze:
brisanje je jedina radnja u cijelom sustavu koja trajno uklanja podatak.

Brisanje retka iz Excela nije samo brisanje retka. Formule u bloku s
ukupnim iznosima pokazuju na redove ispod, a openpyxl ih pri brisanju ne
pomiče - 'Total Income' bi ostao =SUM(L10:L29) u ćeliji L29, dakle kružna
referenca, a PDV i provizije bi gledali u krive retke. Zato se nakon
brisanja formule ispisuju iznova, PDV vraća kakav je bio, a polaznici
prenumeriraju.

Pokretanje:
    python3 otkazi_polaznika.py --tablica "Modul 1&2 Split.xlsx" \\
        --email ana@primjer.hr                 # samo pokaži
    python3 otkazi_polaznika.py --tablica "..." --email ... --stvarno
    python3 otkazi_polaznika.py --tablica "..." --ime "Ana Anić" --stvarno
"""

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import tablice
from send_reminders import COL_EMAIL, COL_FIRST_NAME, COL_LAST_NAME
from zoho_to_excel import (
    FIRST_PARTICIPANT_ROW,
    find_totals_row,
    load_config,
    with_retry,
    write_totals_formulas,
)

COL_UPLATA = 12
ZAPISNIK = Path(__file__).with_name("otkazani.json")
MAPA_KOPIJA = "Otkazani"


def nadji_tablicu(output_dir: Path, naziv: str) -> Path:
    for put in tablice.nadji(output_dir, javi=lambda _: None):
        if put.name == naziv:
            return put
    sys.exit(f"Nema tablice {naziv!r} u {output_dir}.")


def nadji_red(ws, totals_row: int, email: str, ime: str):
    """(broj retka, podaci) ili izlaz s objašnjenjem. Adresa ima prednost -
    jedinstvena je; ime se koristi za polaznika upisanog bez adrese."""
    pogodci = []
    for row in range(FIRST_PARTICIPANT_ROW, totals_row):
        r_ime = str(ws.cell(row=row, column=COL_FIRST_NAME).value or "").strip()
        r_prezime = str(ws.cell(row=row, column=COL_LAST_NAME).value or "").strip()
        r_email = str(ws.cell(row=row, column=COL_EMAIL).value or "").strip()
        if not r_ime and not r_prezime:
            continue
        puno_ime = f"{r_ime} {r_prezime}".strip()
        if email:
            if r_email.lower() == email.lower():
                pogodci.append((row, puno_ime, r_email))
        elif ime and puno_ime.lower() == ime.strip().lower():
            pogodci.append((row, puno_ime, r_email))

    trazeno = email or ime
    if not pogodci:
        sys.exit(f"U tablici nema nikoga pod {trazeno!r}. "
                 f"Možda je tablica u međuvremenu promijenjena.")
    if len(pogodci) > 1:
        redovi = ", ".join(str(p[0]) for p in pogodci)
        sys.exit(f"Pod {trazeno!r} ima više polaznika (redovi {redovi}). "
                 f"Ne diram tablicu - razriješi to ručno.")
    return pogodci[0]


def zabiljezi(podaci: dict) -> None:
    zapisi = []
    if ZAPISNIK.exists():
        try:
            with open(ZAPISNIK, encoding="utf-8") as f:
                zapisi = json.load(f)
        except (ValueError, OSError):
            zapisi = []
    zapisi.append(podaci)
    with open(ZAPISNIK, "w", encoding="utf-8") as f:
        json.dump(zapisi, f, ensure_ascii=False, indent=1)


def otkazi(put: Path, email: str, ime: str, stvarno: bool) -> int:
    wb = with_retry(lambda: load_workbook(put), retry_on=(PermissionError, OSError))
    if "podaci" not in wb.sheetnames:
        sys.exit(f"{put.name} nema list 'podaci'.")
    ws = wb["podaci"]

    totals_row = find_totals_row(ws)
    red, puno_ime, adresa = nadji_red(ws, totals_row, email, ime)
    uplaceno = ws.cell(row=red, column=COL_UPLATA).value or 0

    tecaj = f"{ws['C4'].value or ''} / {ws['C5'].value or ''}".strip(" /")
    print(f"{put.name} ({tecaj})")
    print(f"  redak {red}: {puno_ime}" + (f" <{adresa}>" if adresa else " (bez adrese)"))
    print(f"  uplaćeno: {uplaceno}")

    if not stvarno:
        print("\nPregled - ništa nije promijenjeno. Za stvarno brisanje dodaj --stvarno")
        return 0

    try:
        float_uplaceno = float(uplaceno)
    except (TypeError, ValueError):
        float_uplaceno = 0.0
    if float_uplaceno > 0:
        print(f"  [UPOZORENJE] Ovaj polaznik ima uplatu od {uplaceno}. "
              f"Brisanjem retka nestaje i taj zapis iz tablice - povrat novca "
              f"riješi odvojeno.")

    # 1. kopija cijele tablice, prije ijedne izmjene
    kopije = put.parent / MAPA_KOPIJA
    kopije.mkdir(exist_ok=True)
    kopija = kopije / f"{put.stem} prije otkaza {datetime.now():%Y-%m-%d %H-%M-%S}.xlsx"
    shutil.copy2(put, kopija)
    print(f"  kopija: {kopija.name}")

    # 2. brisanje, pa popravak onoga što openpyxl ne popravlja sam
    pdv = ws.cell(row=totals_row + 2, column=11).value
    ws.delete_rows(red)
    novi_totals = find_totals_row(ws)
    write_totals_formulas(ws, novi_totals)
    if pdv is not None:
        ws.cell(row=novi_totals + 2, column=11, value=pdv)
    for i, r in enumerate(range(FIRST_PARTICIPANT_ROW, novi_totals), start=1):
        ws.cell(row=r, column=1, value=i)

    wb.save(put)

    zabiljezi({
        "kada": datetime.now().isoformat(timespec="seconds"),
        "tablica": put.name,
        "tecaj": tecaj,
        "ime": puno_ime,
        "email": adresa,
        "uplaceno": float_uplaceno,
        "kopija": str(kopija),
    })
    print(f"  izbrisan iz tablice, zapisano u {ZAPISNIK.name}")
    return 1


def main():
    parser = argparse.ArgumentParser(
        description="Izbaci polaznika iz tablice tečaja (kad otkaže dolazak).")
    parser.add_argument("--tablica", required=True, metavar="NAZIV.xlsx",
                        help="Točan naziv datoteke tablice")
    parser.add_argument("--email", default="", help="Adresa polaznika")
    parser.add_argument("--ime", default="", help="Ime i prezime (za polaznika bez adrese)")
    parser.add_argument("--stvarno", action="store_true",
                        help="Stvarno izbriši; bez ovoga se samo prikaže tko bi otišao")
    args = parser.parse_args()

    if not args.email and not args.ime:
        sys.exit("Treba --email ili --ime.")

    config = load_config()
    output_dir = Path(config["output_dir"])
    if not output_dir.exists():
        sys.exit(f"Ne postoji folder s tablicama: {output_dir}")

    put = nadji_tablicu(output_dir, args.tablica)
    broj = otkazi(put, args.email, args.ime, args.stvarno)
    if args.stvarno:
        print(f"\nGotovo. Izbrisano {broj}.")


if __name__ == "__main__":
    main()
