"""
Pošalje polaznicima njihove potvrdnice - na tvoju naredbu, ne sam.

Potvrdnice prvo napravi napravi_certifikate.py; ova skripta ih uzme iz
mape tečaja, iz tablice pročita email adrese i svakome pošalje njegovu, i
samo njegovu, potvrdnicu u privitku.

NIŠTA SE NE ŠALJE DOK NE DODAŠ --posalji. Bez toga skripta samo pokaže
kome bi što poslala - tako uvijek prvo vidiš popis.

    python3 posalji_certifikate.py                    # pokaži kome bi poslao
    python3 posalji_certifikate.py --posalji          # pošalji
    python3 posalji_certifikate.py --posalji --tecaj Split   # samo taj tečaj

Tekst maila stoji u dokumentu 'mail uz certifikat.docx' uz Excel tablice
(isti princip kao podsjetnici); ako ga nema, koristi se ugrađeni tekst.
Poveznica na termine sljedećih tečajeva se upisuje kao {link_termini}, a
sama adresa stoji u config.json pod "course_dates_url".

Zapamti kome je poslano (poslani_certifikati.json), pa nitko ne dobiva
potvrdnicu dvaput, koliko god puta skriptu pokrenuo.
"""

import argparse
import json
import mimetypes
import smtplib
import sys
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

import predlosci
import tablice
from napravi_certifikate import ime_datoteke, ime_mape
from send_reminders import (
    FIRST_PARTICIPANT_ROW,
    course_info,
    datumi_rijecima,
    find_totals_row,
    parse_course_dates,
)
from zoho_to_excel import load_config, with_retry

STATE_PATH = Path(__file__).with_name("poslani_certifikati.json")

ZADANI_NASLOV = "Vaša potvrdnica - Emmett tehnika {course_code}, {location}"

ZADANI_TEKST = """Poštovani/a {first_name} {last_name},

hvala vam na sudjelovanju na tečaju Emmett tehnike za ljude - {course_code}, koji smo održali {datumi_rijecima} u mjestu {location}.

U privitku se nalazi vaša potvrdnica o završenom tečaju (Statement of Attainment), izdana od strane EMMETT Technique International Headquarters. Preporučujem da je spremite; trebat će vam kao dokaz o završenom modulu pri upisu na sljedeće stupnjeve edukacije.

Bilo mi je zadovoljstvo raditi s vama ovaj vikend. Ako vam tijekom rada iskrsne neko pitanje o naučenim tehnikama, slobodno mi se javite.

Termine sljedećih modula i prijavnicu možete pronaći ovdje:
{link_termini}

Želim vam puno uspjeha u radu.

Lijep pozdrav,
{instructor_name}
EMMETT Hrvatska
"""

PDF_MIME = "application/pdf"


def load_state() -> dict:
    if STATE_PATH.exists():
        with open(STATE_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_state(state: dict) -> None:
    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2, sort_keys=True)


def citaj_polaznike(ws, totals_row: int) -> list:
    """[{ime, email}] - ime iz stupaca B i C, adresa iz H."""
    polaznici = []
    for row in range(FIRST_PARTICIPANT_ROW, totals_row):
        ime = ws.cell(row=row, column=2).value
        if not ime:
            continue
        prezime = ws.cell(row=row, column=3).value or ""
        email = ws.cell(row=row, column=8).value
        polaznici.append({
            "ime": f"{str(ime).strip()} {str(prezime).strip()}".strip(),
            "first_name": str(ime).strip(),
            "last_name": str(prezime).strip(),
            "email": str(email).strip() if email else "",
        })
    return polaznici


def posalji_jedan(config: dict, primatelj: str, naslov: str, tijelo: str,
                  pdf: Path) -> None:
    msg = EmailMessage()
    msg["Subject"] = naslov
    msg["From"] = config["zoho_email"]
    msg["To"] = primatelj
    msg["Bcc"] = config["zoho_email"]      # kopija tebi, da imaš trag u Sentu
    msg.set_content(tijelo)

    tip, _ = mimetypes.guess_type(pdf.name)
    glavni, _, pod = (tip or PDF_MIME).partition("/")
    msg.add_attachment(pdf.read_bytes(), maintype=glavni,
                       subtype=pod or "pdf", filename=pdf.name)

    def _posalji():
        with smtplib.SMTP_SSL(config["smtp_host"], config.get("smtp_port", 465)) as smtp:
            smtp.login(config["zoho_email"], config["zoho_app_password"])
            smtp.send_message(msg)

    with_retry(_posalji)


def tekst_iz_datoteka(naslov_put: str, tijelo_put: str):
    """Naslov i tijelo koje je zadala web stranica, umjesto predloška uz
    tablice. Oznake tipa {first_name} ostaju i ovdje - poruka ide svakom
    polazniku s njegovim imenom, pa se popunjavaju po osobi."""
    def procitaj(put_tekst, sto):
        put = Path(put_tekst)
        if not put.exists():
            sys.exit(f"Nema datoteke s {sto}: {put}")
        tekst = put.read_text(encoding="utf-8").strip()
        if not tekst:
            sys.exit(f"Datoteka s {sto} je prazna: {put}")
        return tekst
    return (procitaj(naslov_put, "naslovom"), procitaj(tijelo_put, "tijelom poruke"),
            "tekst zadan sa stranice")


def obradi(xlsx_path: Path, config: dict, state: dict, posalji: bool,
           zadani_tekst=None) -> int:
    """Vrati broj poslanih (ili, u pregledu, spremnih) potvrdnica."""
    wb = with_retry(lambda: load_workbook(xlsx_path), retry_on=(PermissionError, OSError))
    if "podaci" not in wb.sheetnames:
        return 0
    ws = wb["podaci"]

    info = course_info(ws, config)
    _, kraj = parse_course_dates(info["dates"])
    if kraj is None:
        return 0

    mapa = xlsx_path.parent / "Certifikati" / ime_mape(info, kraj)
    if not mapa.is_dir():
        return 0   # za ovaj tečaj potvrdnice još nisu napravljene

    polaznici = citaj_polaznike(ws, find_totals_row(ws))
    vec_poslano = {a.lower() for a in state.get(xlsx_path.name, {}).get("poslano", [])}

    if zadani_tekst:
        naslov_pred, tijelo_pred, izvor = zadani_tekst
    else:
        naslov_pred, tijelo_pred, izvor = predlosci.dohvati(
            xlsx_path.parent, "mail uz certifikat", config=config,
            kljuc_naslov="certificate_mail_subject", kljuc_tijela="certificate_mail_body",
            zadani_naslov=ZADANI_NASLOV, zadano_tijelo=ZADANI_TEKST)

    print(f"\n{info['course_code']} / {info['location']} ({info['dates']})")
    print(f"  mapa:  {mapa}")
    print(f"  tekst: {izvor}")

    cekaju, problemi = [], []
    for p in polaznici:
        pdf = mapa / f"{ime_datoteke(p['ime'])}.pdf"
        if not pdf.exists():
            problemi.append(f"{p['ime']}: nema potvrdnice ({pdf.name})")
        elif "@" not in p["email"]:
            problemi.append(f"{p['ime']}: nema email adrese u tablici")
        elif p["email"].lower() in vec_poslano:
            pass   # već dobio
        else:
            cekaju.append((p, pdf))

    for x in problemi:
        print(f"  [!] {x}")
    if not cekaju:
        print(f"  Nema kome slati (već poslano: {len(vec_poslano)}).")
        return 0

    varijable = dict(info,
                     datumi_rijecima=datumi_rijecima(info["dates"]),
                     link_termini=config.get("course_dates_url", ""))

    if not posalji:
        for p, pdf in cekaju:
            print(f"  - {p['ime']} <{p['email']}>  ->  {pdf.name}")
        prvi = cekaju[0][0]
        print("  --- poruka ---")
        v = dict(varijable, **{k: prvi[k] for k in ("first_name", "last_name")})
        print("  Naslov:", predlosci.popuni(naslov_pred, v, izvor))
        for redak in predlosci.popuni(tijelo_pred, v, izvor).splitlines():
            print("  " + redak)
        return len(cekaju)

    poslano = set(vec_poslano)
    for p, pdf in cekaju:
        v = dict(varijable, **{k: p[k] for k in ("first_name", "last_name")})
        try:
            posalji_jedan(config, p["email"],
                          predlosci.popuni(naslov_pred, v, izvor),
                          predlosci.popuni(tijelo_pred, v, izvor), pdf)
        except Exception as e:
            print(f"  [!] {p['ime']} <{p['email']}>: {e}")
            continue
        poslano.add(p["email"].lower())
        print(f"  Poslano: {p['ime']} <{p['email']}>")
        # spremaj nakon svake poruke - ako slanje pukne na pola, sljedeće
        # pokretanje nastavlja bez dupliranja
        state[xlsx_path.name] = {
            "tecaj": f"{info['course_code']} / {info['location']}",
            "datumi": info["dates"],
            "poslano": sorted(poslano),
        }
        save_state(state)

    return len(poslano) - len(vec_poslano)


def main():
    parser = argparse.ArgumentParser(
        description="Pošalje polaznicima njihove potvrdnice (na tvoju naredbu).")
    parser.add_argument("--posalji", action="store_true",
                        help="Stvarno pošalji; bez ovoga se samo prikaže popis")
    parser.add_argument("--tecaj", metavar="DIO NAZIVA",
                        help="Samo tečajevi čiji naziv tablice sadrži ovaj tekst")
    parser.add_argument("--naslov-datoteka", dest="naslov_datoteka",
                        help="Datoteka s naslovom, umjesto predloška uz tablice")
    parser.add_argument("--tijelo-datoteka", dest="tijelo_datoteka",
                        help="Datoteka s tijelom poruke, umjesto predloška")
    args = parser.parse_args()

    if bool(args.naslov_datoteka) != bool(args.tijelo_datoteka):
        sys.exit("--naslov-datoteka i --tijelo-datoteka idu zajedno.")
    zadani_tekst = (tekst_iz_datoteka(args.naslov_datoteka, args.tijelo_datoteka)
                    if args.naslov_datoteka else None)

    config = load_config()
    output_dir = Path(config["output_dir"])
    if not output_dir.exists():
        sys.exit(f"Ne postoji folder s tablicama: {output_dir}")

    datoteke = tablice.nadji(output_dir)
    if args.tecaj:
        trazeno = predlosci.bez_dijakritika(args.tecaj)
        datoteke = [p for p in datoteke if trazeno in predlosci.bez_dijakritika(p.name)]
        if not datoteke:
            sys.exit(f"Nijedna tablica u nazivu nema {args.tecaj!r}.")

    state = load_state()
    ukupno = 0
    for put in datoteke:
        try:
            ukupno += obradi(put, config, state, args.posalji, zadani_tekst)
        except predlosci.GreskaPredloska as e:
            print(f"[!] {put.name}: {e}")

    print()
    if args.posalji:
        print(f"Gotovo. Poslano {ukupno} potvrdnica.")
    elif ukupno:
        print(f"Spremno za slanje: {ukupno} potvrdnica.")
        print("Za stvarno slanje pokreni istu naredbu s  --posalji")
    else:
        print("Nema nijedne potvrdnice za poslati.")


if __name__ == "__main__":
    main()
