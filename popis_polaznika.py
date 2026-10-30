"""
Napravi PDF s popisom polaznika jednog tečaja - samo imena i prezimena.

Na popisu nema adresa, telefona ni iznosa. To je namjerno: popis se nosi na
tečaj i pokazuje drugima, pa na njemu ne smije biti ničega što se tiče samo
polaznika i tebe.

U zaglavlju piše koji je modul, kada, gdje i tko ga vodi - sve iz same
tablice.

PDF nastaje uz tablice, pa ga Dropbox odnese i na mobitel:

    <output_dir>/Popisi/<Mjesto> <Kod tečaja> <Godina>.pdf

Pokretanje:
    python3 popis_polaznika.py --tecaj Split
    python3 popis_polaznika.py --tecaj Split --pregled   # samo ispiši imena

Pokreće ga i tipka "Popis polaznika" na stranici tečaja.
"""

import argparse
import base64
import html
import sys
from pathlib import Path

from openpyxl import load_workbook

import predlosci
import tablice
from napravi_certifikate import (
    citaj_polaznike,
    iscrtaj,
    ime_datoteke,
    ime_mape,
    nadji_chrome,
)
from send_reminders import course_info, find_totals_row, parse_course_dates
from zoho_to_excel import load_config, with_retry

# Wordmark stoji uz tablice, u mapi 'Marka' - font mu je komercijalan, pa ne
# može u javni repozitorij. Ako ga nema, popis ide bez njega.
WORDMARK = "wordmark-navy.svg"

PREDLOZAK = """<!DOCTYPE html>
<html lang="hr">
<head>
<meta charset="utf-8">
<title>{naslov}</title>
<style>
  @page {{ size: A4; margin: 18mm 16mm 16mm 16mm; }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: Figtree, "Liberation Sans", Arial, Helvetica, sans-serif;
    color: #1b2444;
    -webkit-print-color-adjust: exact; print-color-adjust: exact;
  }}
  .marka {{ height: 26px; margin-bottom: 18px; }}
  h1 {{ font-size: 22pt; margin: 0 0 2mm; letter-spacing: .2pt; }}
  .podaci {{ font-size: 11pt; color: #41507e; margin: 0 0 1mm; }}
  .crta {{ border-bottom: 1.4pt solid #1b2444; margin: 5mm 0 4mm; }}
  .koliko {{ font-size: 10pt; color: #6b779f; margin: 0 0 4mm; }}
  ol {{ margin: 0; padding: 0; list-style: none; counter-reset: redni; }}
  li {{
    counter-increment: redni;
    display: flex; align-items: baseline; gap: 6mm;
    padding: 2.6mm 0; border-bottom: .4pt solid #d7dceb;
    font-size: 12pt;
    break-inside: avoid;
  }}
  li::before {{
    content: counter(redni) ".";
    width: 10mm; text-align: right;
    font-size: 10pt; color: #6b779f; font-variant-numeric: tabular-nums;
  }}
  .nema {{ font-size: 12pt; color: #6b779f; }}
</style>
</head>
<body>
{marka}
<h1>{kod}</h1>
<p class="podaci">{datumi} · {mjesto}</p>
<p class="podaci">Instruktor: {instruktor}</p>
<div class="crta"></div>
<p class="koliko">{koliko}</p>
{popis}
</body>
</html>
"""


def oblik_broja(koliko: int) -> str:
    return "1 polaznik" if koliko == 1 else f"{koliko} polaznika"


def sastavi(info: dict, polaznici: list, marka: Path) -> str:
    if polaznici:
        redci = "\n".join(f"  <li>{html.escape(ime)}</li>" for ime in polaznici)
        popis = f"<ol>\n{redci}\n</ol>"
    else:
        popis = '<p class="nema">U tablici još nema nijednog polaznika.</p>'

    znak = ""
    if marka.is_file():
        # wordmark ide ugrađen, a ne kao vanjska datoteka: Chrome iscrtava
        # HTML iz privremene datoteke uz PDF, odakle putanja do Marke ne vodi
        # nikamo.
        ugradeno = base64.b64encode(marka.read_bytes()).decode()
        znak = (f'<img class="marka" alt="EMMETT" '
                f'src="data:image/svg+xml;base64,{ugradeno}">')

    return PREDLOZAK.format(
        naslov=html.escape(f"{info['course_code']} - {info['location']}"),
        marka=znak,
        kod=html.escape(info["course_code"]),
        datumi=html.escape(info["dates"]),
        mjesto=html.escape(info["location"]),
        instruktor=html.escape(info["instructor_name"]),
        koliko=oblik_broja(len(polaznici)),
        popis=popis)


def obradi(xlsx_path: Path, config: dict, chrome: str, dry_run: bool) -> Path:
    wb = with_retry(lambda: load_workbook(xlsx_path),
                    retry_on=(PermissionError, OSError))
    if "podaci" not in wb.sheetnames:
        print(f"[!] {xlsx_path.name}: nema list 'podaci', preskačem.")
        return None
    ws = wb["podaci"]

    info = course_info(ws, config)
    polaznici = citaj_polaznike(ws, find_totals_row(ws))
    _, kraj = parse_course_dates(info["dates"])
    godina = kraj.year if kraj else ""

    print(f"\n{info['course_code']} / {info['location']} ({info['dates']}) - "
          f"{len(polaznici)} polaznika")
    print(f"  instruktor: {info['instructor_name']}")
    if dry_run:
        for redni, ime in enumerate(polaznici, start=1):
            print(f"    {redni}. {ime}")
        return None

    mapa = Path(config["output_dir"]) / "Popisi"
    mapa.mkdir(parents=True, exist_ok=True)
    naziv = ime_datoteke(f"{info['location']} {info['course_code']} {godina}".strip())
    pdf_put = mapa / f"{naziv}.pdf"

    html_tekst = sastavi(info, polaznici, Path(config["output_dir"]) / "Marka" / WORDMARK)
    iscrtaj(chrome, html_tekst, pdf_put, config.get("chrome_args", []))
    print(f"  {pdf_put}")
    return pdf_put


def main():
    parser = argparse.ArgumentParser(
        description="PDF s popisom polaznika - samo imena i prezimena.")
    parser.add_argument("--tecaj", metavar="DIO NAZIVA",
                        help="Samo tečajevi čiji naziv tablice sadrži ovaj tekst")
    parser.add_argument("--pregled", action="store_true",
                        help="Samo ispiši imena, bez pravljenja PDF-a")
    args = parser.parse_args()

    config = load_config()
    output_dir = Path(config["output_dir"])
    if not output_dir.exists():
        sys.exit(f"Ne postoji folder s tablicama: {output_dir}")

    datoteke = tablice.nadji(output_dir)
    if args.tecaj:
        trazeno = predlosci.bez_dijakritika(args.tecaj)
        datoteke = [p for p in datoteke
                    if trazeno in predlosci.bez_dijakritika(p.name)]
        if not datoteke:
            sys.exit(f"Nijedna tablica u nazivu nema {args.tecaj!r}.")

    chrome = "" if args.pregled else nadji_chrome(config)
    napravljeno = 0
    for put in datoteke:
        if obradi(put, config, chrome, args.pregled):
            napravljeno += 1

    print()
    if args.pregled:
        print("Pregled gotov - nijedan PDF nije napravljen.")
    else:
        print(f"Gotovo. Napravljeno {napravljeno} popisa.")


if __name__ == "__main__":
    main()
