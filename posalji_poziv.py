"""
Pozovi polaznike prethodnog modula da se prijave na tečaj.

Poruka ide svakome zasebno, sa svojim imenom - nitko ne vidi tuđe adrese.
Koga pozvati odlučuje se na stranici (vidi pozivi.py); ovamo stiže gotov
popis, pa skripta ne pogađa ništa.

Tko je jednom pozvan na određeni tečaj, neće biti pozvan opet - to pamti
poslani_pozivi.json, isto načelo kao kod potvrdnica.

Pokretanje (primatelji i tekst u datotekama, kao i kod ostalih poruka):

    python3 posalji_poziv.py --tecaj "Modul 3 Zagreb.xlsx" \\
        --primatelji ljudi.json \\
        --naslov-datoteka naslov.txt --tijelo-datoteka tijelo.txt
    python3 posalji_poziv.py ... --posalji
"""

import argparse
import json
import smtplib
import sys
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

import pozivi
import predlosci
import tablice
from send_reminders import course_info
from zoho_to_excel import load_config, with_retry

STATE_PATH = Path(__file__).with_name("poslani_pozivi.json")


def ucitaj_stanje() -> dict:
    if STATE_PATH.exists():
        try:
            with open(STATE_PATH, encoding="utf-8") as f:
                return json.load(f)
        except (ValueError, OSError):
            return {}
    return {}


def spremi_stanje(stanje: dict) -> None:
    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(stanje, f, ensure_ascii=False, indent=2, sort_keys=True)


def procitaj_tekst(put_tekst: str, sto: str) -> str:
    put = Path(put_tekst)
    if not put.exists():
        sys.exit(f"Nema datoteke s {sto}: {put}")
    tekst = put.read_text(encoding="utf-8").strip()
    if not tekst:
        sys.exit(f"Datoteka s {sto} je prazna: {put}")
    return tekst


def procitaj_primatelje(put_tekst: str) -> list:
    put = Path(put_tekst)
    if not put.exists():
        sys.exit(f"Nema popisa primatelja: {put}")
    try:
        with open(put, encoding="utf-8") as f:
            ljudi = json.load(f)
    except (ValueError, OSError) as e:
        sys.exit(f"Ne mogu pročitati popis primatelja: {e}")
    if not isinstance(ljudi, list) or not ljudi:
        sys.exit("Popis primatelja je prazan.")
    for osoba in ljudi:
        if not isinstance(osoba, dict) or "@" not in str(osoba.get("email", "")):
            sys.exit(f"Neispravan zapis u popisu primatelja: {osoba!r}")
    return ljudi


def nadji_tablicu(output_dir: Path, naziv: str) -> Path:
    for put in tablice.nadji(output_dir, javi=lambda _: None):
        if put.name == naziv:
            return put
    sys.exit(f"Nema tablice {naziv!r} u {output_dir}.")


def posalji_jedan(config: dict, kome: str, naslov: str, tijelo: str) -> None:
    msg = EmailMessage()
    msg["Subject"] = naslov
    msg["From"] = config["zoho_email"]
    msg["To"] = kome
    msg["Bcc"] = config["zoho_email"]
    msg.set_content(tijelo)

    def _posalji():
        with smtplib.SMTP_SSL(config["smtp_host"], config.get("smtp_port", 465)) as smtp:
            smtp.login(config["zoho_email"], config["zoho_app_password"])
            smtp.send_message(msg)

    with_retry(_posalji)


def main():
    parser = argparse.ArgumentParser(
        description="Pozovi polaznike prethodnog modula na tečaj.")
    parser.add_argument("--tecaj", required=True, metavar="NAZIV.xlsx")
    parser.add_argument("--primatelji", required=True,
                        help="JSON s popisom [{ime, email}]")
    parser.add_argument("--naslov-datoteka", required=True, dest="naslov_datoteka")
    parser.add_argument("--tijelo-datoteka", required=True, dest="tijelo_datoteka")
    parser.add_argument("--posalji", action="store_true",
                        help="Stvarno pošalji; bez ovoga je samo pregled")
    args = parser.parse_args()

    config = load_config()
    output_dir = Path(config["output_dir"])
    put = nadji_tablicu(output_dir, args.tecaj)

    wb = with_retry(lambda: load_workbook(put), retry_on=(PermissionError, OSError))
    if "podaci" not in wb.sheetnames:
        sys.exit(f"{put.name} nema list 'podaci'.")
    info = course_info(wb["podaci"], config)

    naslov_pred = procitaj_tekst(args.naslov_datoteka, "naslovom")
    tijelo_pred = procitaj_tekst(args.tijelo_datoteka, "tijelom poruke")
    ljudi = procitaj_primatelje(args.primatelji)

    stanje = ucitaj_stanje()
    vec = {a.lower() for a in stanje.get(put.name, {}).get("poslano", [])}
    cekaju = [o for o in ljudi if o["email"].lower() not in vec]

    print(f"{info['course_code']} / {info['location']} ({info['dates']})")
    print(f"  na popisu: {len(ljudi)}, već pozvano: {len(ljudi) - len(cekaju)}")
    if not cekaju:
        print("  Nema kome slati.")
        return

    osnova = {
        "course_code": info["course_code"],
        "location": info["location"],
        "dates": info["dates"],
        "instructor_name": info["instructor_name"],
        "link_termini": config.get("course_dates_url", ""),
        "prethodni_modul": pozivi.naziv_prethodnog(config, info["course_code"])
                           or "prethodni modul",
    }

    if not args.posalji:
        for o in cekaju:
            print(f"  - {o['ime']} <{o['email']}>")
        prvi = cekaju[0]
        v = dict(osnova, **imena(prvi["ime"]))
        print("  --- poruka ---")
        print("  Naslov:", predlosci.popuni(naslov_pred, v, "zadani tekst"))
        for redak in predlosci.popuni(tijelo_pred, v, "zadani tekst").splitlines():
            print("  " + redak)
        print("\nPregled - ništa nije poslano. Za slanje dodaj --posalji")
        return

    poslano = set(stanje.get(put.name, {}).get("poslano", []))
    for o in cekaju:
        v = dict(osnova, **imena(o["ime"]))
        try:
            posalji_jedan(config, o["email"],
                          predlosci.popuni(naslov_pred, v, "zadani tekst"),
                          predlosci.popuni(tijelo_pred, v, "zadani tekst"))
        except Exception as e:
            print(f"  [!] {o['ime']} <{o['email']}>: {e}")
            continue
        poslano.add(o["email"])
        print(f"  Poslano: {o['ime']} <{o['email']}>")
        # spremaj nakon svake poruke - ako slanje pukne na pola, sljedeće
        # pokretanje nastavlja bez dupliranja
        stanje[put.name] = {
            "tecaj": f"{info['course_code']} / {info['location']}",
            "datumi": info["dates"],
            "poslano": sorted(poslano),
        }
        spremi_stanje(stanje)

    print(f"\nGotovo. Poslano {len(poslano) - len(vec)} poziva.")


def imena(puno_ime: str) -> dict:
    dijelovi = (puno_ime or "").split()
    return {"first_name": dijelovi[0] if dijelovi else "",
            "last_name": " ".join(dijelovi[1:]),
            "ime": puno_ime or ""}


if __name__ == "__main__":
    main()
