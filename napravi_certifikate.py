"""
Dan nakon tečaja napravi EMMETT potvrdnice (Statement of Attainment) za
svakog polaznika iz tablice.

Iz tablice se čitaju mjesto, kod tečaja, datumi i imena polaznika, pa se
za svakoga iscrta PDF prema odgovarajućoj potvrdnici:

    Modul 1&2            -> potvrdnica modul 1 i 2.html
    Modul 3, 4, 5, 6     -> potvrdnica modul 3-6.html (modul se upisuje sam)

Oba predloška stoje uz Excel tablice (output_dir) - isti princip kao
'lokacija split.docx'. Izvedeni su iz potvrdnica u Claude Designu; kad ih
ondje promijeniš, treba ih ponovno izvesti.

PDF-ovi se slažu u:

    <output_dir>/Certifikati/<Mjesto> <Kod tečaja> <Godina>/Ime Prezime.pdf

Ništa se ne šalje mailom - datoteke samo nastanu i čekaju te.

Za iscrtavanje se koristi Chrome (ili Chromium / Edge) u pozadini; put se
po potrebi zadaje s "chrome_path" u config.json.

Pokretanje:
    python3 napravi_certifikate.py --pregled     # samo pokaži što bi napravio
    python3 napravi_certifikate.py               # napravi PDF-ove
    python3 napravi_certifikate.py --ponovno     # napravi i one koji već postoje

Iz crona, jednom dnevno uz posalji_tablicu.py.
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import tablice
from send_reminders import (
    FIRST_PARTICIPANT_ROW,
    course_info,
    find_totals_row,
    parse_course_dates,
)
from zoho_to_excel import load_config, with_retry

STATE_PATH = Path(__file__).with_name("napravljeni_certifikati.json")

MJESECI = ["January", "February", "March", "April", "May", "June",
           "July", "August", "September", "October", "November", "December"]

# Koja potvrdnica za koji tečaj. Ključ je kod tečaja iz tablice (C4),
# vrijednost je (naziv datoteke predloška, oznaka modula ili None).
PREDLOSCI = {
    "modul 1&2": ("potvrdnica modul 1 i 2", None),
    "modul 3": ("potvrdnica modul 3-6", "3"),
    "modul 4": ("potvrdnica modul 3-6", "4"),
    "modul 5": ("potvrdnica modul 3-6", "5"),
    "modul 6": ("potvrdnica modul 3-6", "6"),
}

# Chrome iscrtava PDF. Prvi koji postoji se koristi.
PUTEVI_CHROMEA = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
    "/snap/bin/chromium",
]


def nadji_chrome(config: dict) -> str:
    zadano = config.get("chrome_path")
    if zadano:
        if Path(zadano).exists():
            return zadano
        sys.exit(f"chrome_path iz config.json ne postoji: {zadano}")
    for put in PUTEVI_CHROMEA:
        if Path(put).exists():
            return put
    nadjen = shutil.which("google-chrome") or shutil.which("chromium")
    if nadjen:
        return nadjen
    sys.exit("Ne mogu naći Chrome - bez njega se PDF ne može napraviti.\n"
             "Instaliraj Google Chrome, ili upiši put u config.json kao "
             '"chrome_path": "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"')


def nastavak_dana(dan: int) -> str:
    """4 -> '4th', 1 -> '1st', 22 -> '22nd', 11 -> '11th'."""
    if 11 <= dan % 100 <= 13:
        return f"{dan}th"
    return f"{dan}{ {1: 'st', 2: 'nd', 3: 'rd'}.get(dan % 10, 'th') }"


def predlozak_za(course_code: str, mapa: Path):
    """(put do predloška, oznaka modula) ili (None, razlog zašto ne ide)."""
    kljuc = " ".join(str(course_code).lower().split())
    if kljuc not in PREDLOSCI:
        return None, (f"za tečaj {course_code!r} nema potvrdnice "
                      f"(postoje samo za Modul 1&2 i Modul 3-6)")
    naziv, modul = PREDLOSCI[kljuc]
    put = mapa / f"{naziv}.html"
    if not put.exists():
        return None, f"nedostaje predložak '{put.name}' u {mapa}"
    return (put, modul), None


def ime_mape(info: dict, kraj: date) -> str:
    """'Split Modul 1&2 2026' - mjesto, kod tečaja i godina."""
    naziv = f"{info['location']} {info['course_code']} {kraj.year}".strip()
    return re.sub(r'[\\/:*?"<>|]', "", naziv).strip()


def citaj_polaznike(ws, totals_row: int) -> list:
    """Imena izravno iz redaka - potvrdnica ne treba email, pa se uzima i
    polaznik upisan ručno, bez adrese."""
    polaznici = []
    for row in range(FIRST_PARTICIPANT_ROW, totals_row):
        ime = ws.cell(row=row, column=2).value
        if not ime:
            continue
        prezime = ws.cell(row=row, column=3).value or ""
        polaznici.append(f"{str(ime).strip()} {str(prezime).strip()}".strip())
    return polaznici


def ime_datoteke(ime: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "", ime).strip() or "Bez imena"


def iscrtaj(chrome: str, html: str, pdf_put: Path, dodatno=()) -> None:
    """Chrome iscrta HTML u PDF. HTML ide u privremenu datoteku jer se
    predložak puni po polazniku. `dodatno` su zastavice iz config.json
    ("chrome_args") - na Macu ne trebaju, služe za neobična okruženja.

    Privremena datoteka namjerno nastaje UZ PDF, a ne u /tmp: Chromium
    instaliran kao snap ima vlastiti /tmp i onaj pravi ne vidi, pa bi odande
    pročitao prazno. Mapa odredišta mu je dostupna jer u nju i piše."""
    izvor = pdf_put.with_suffix(".tmp.html")
    try:
        izvor.write_text(html, encoding="utf-8")
        rezultat = subprocess.run(
            [chrome, "--headless", "--disable-gpu", "--no-pdf-header-footer",
             *dodatno, f"--print-to-pdf={pdf_put}", izvor.as_uri()],
            capture_output=True, text=True, timeout=180)
    finally:
        izvor.unlink(missing_ok=True)
    if not pdf_put.exists() or pdf_put.stat().st_size == 0:
        poruka = (rezultat.stderr or rezultat.stdout or "").strip().splitlines()
        kraj = poruka[-1] if poruka else f"Chrome je vratio {rezultat.returncode}"
        raise RuntimeError(f"PDF nije napravljen: {kraj}")


def load_state() -> dict:
    if STATE_PATH.exists():
        with open(STATE_PATH, encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_state(state: dict) -> None:
    with open(STATE_PATH, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2, sort_keys=True)


def obradi(xlsx_path: Path, config: dict, state: dict, danas: date,
           chrome: str, dry_run: bool, ponovno: bool) -> int:
    """Vrati broj napravljenih (ili, u pregledu, predviđenih) potvrdnica."""
    wb = with_retry(lambda: load_workbook(xlsx_path), retry_on=(PermissionError, OSError))
    if "podaci" not in wb.sheetnames:
        return 0
    ws = wb["podaci"]

    info = course_info(ws, config)
    pocetak, kraj = parse_course_dates(info["dates"])
    if kraj is None:
        print(f"[!] {xlsx_path.name}: ne mogu pročitati datum tečaja "
              f"({info['dates']!r}), preskačem.")
        return 0

    proslo = (danas - kraj).days
    if proslo < config.get("certificate_days_after", 1):
        return 0   # tečaj još traje ili je prerano

    if state.get(xlsx_path.name) and not ponovno:
        return 0   # već napravljeno

    mapa = xlsx_path.parent
    odabir, razlog = predlozak_za(info["course_code"], mapa)
    if odabir is None:
        print(f"[!] {xlsx_path.name}: {razlog} - preskačem.")
        return 0
    predlozak_put, modul = odabir

    polaznici = citaj_polaznike(ws, find_totals_row(ws))
    if not polaznici:
        print(f"[!] {xlsx_path.name}: tečaj je završio, ali u tablici nema "
              f"nijednog polaznika - nemam za koga raditi potvrdnice.")
        return 0

    cilj = mapa / "Certifikati" / ime_mape(info, kraj)
    print(f"\n{info['course_code']} / {info['location']} ({info['dates']}) - "
          f"završio prije {proslo} dana, {len(polaznici)} polaznika")
    print(f"  predložak: {predlozak_put.name}" + (f" (modul {modul})" if modul else ""))
    print(f"  mapa:      {cilj}")

    predlozak = predlozak_put.read_text(encoding="utf-8")
    datum = {"@@DAN@@": nastavak_dana(kraj.day),
             "@@MJESEC@@": MJESECI[kraj.month - 1],
             "@@GODINA@@": str(kraj.year)}

    if dry_run:
        for ime in polaznici:
            print(f"    - {ime_datoteke(ime)}.pdf")
        print(f"  Datum na potvrdnici: {datum['@@DAN@@']} day of "
              f"{datum['@@MJESEC@@']} {datum['@@GODINA@@']}")
        return len(polaznici)

    cilj.mkdir(parents=True, exist_ok=True)
    dodatno = config.get("chrome_args", [])
    napravljeno = []
    for ime in polaznici:
        html = predlozak.replace("@@IME@@", ime)
        for kljuc, vrijednost in datum.items():
            html = html.replace(kljuc, vrijednost)
        if modul:
            html = html.replace("@@MODUL@@", modul)

        pdf_put = cilj / f"{ime_datoteke(ime)}.pdf"
        try:
            iscrtaj(chrome, html, pdf_put, dodatno)
        except Exception as e:
            print(f"    [!] {ime}: {e}")
            continue
        napravljeno.append(ime)
        print(f"    {pdf_put.name}")

    if not napravljeno:
        # ne ostavljaj praznu mapu iza neuspjeha
        try:
            cilj.rmdir()
        except OSError:
            pass
        return 0

    if napravljeno:
        state[xlsx_path.name] = {
            "napravljeno": danas.isoformat(),
            "tecaj": f"{info['course_code']} / {info['location']}",
            "datumi": info["dates"],
            "mapa": str(cilj),
            "polaznici": napravljeno,
        }
        save_state(state)
    return len(napravljeno)


def main():
    parser = argparse.ArgumentParser(
        description="Napravi EMMETT potvrdnice za polaznike nakon tečaja.")
    parser.add_argument("--pregled", action="store_true",
                        help="Samo prikaži što bi napravio, bez pisanja PDF-ova")
    parser.add_argument("--ponovno", action="store_true",
                        help="Napravi potvrdnice i za tečaj koji je već obrađen")
    parser.add_argument("--datum", metavar="GGGG-MM-DD",
                        help="Radi kao da je taj datum (samo uz --pregled)")
    args = parser.parse_args()

    if args.datum and not args.pregled:
        sys.exit("--datum se smije koristiti samo uz --pregled.")

    config = load_config()
    output_dir = Path(config["output_dir"])
    if not output_dir.exists():
        sys.exit(f"Ne postoji folder s tablicama: {output_dir}")

    datoteke = tablice.nadji(output_dir)
    if not datoteke:
        print(f"Nema nijedne .xlsx tablice u {output_dir}.")
        return

    if args.datum:
        try:
            danas = date.fromisoformat(args.datum)
        except ValueError:
            sys.exit(f"Neispravan datum: {args.datum!r} - očekujem oblik 2026-10-05.")
        print(f"Pregled kao da je {danas.strftime('%d.%m.%Y.')}\n")
    else:
        danas = date.today()

    chrome = "" if args.pregled else nadji_chrome(config)
    state = load_state()

    ukupno = 0
    for put in datoteke:
        ukupno += obradi(put, config, state, danas, chrome, args.pregled, args.ponovno)

    print()
    if args.pregled:
        print(f"Pregled gotov ({len(datoteke)} tablica) - ništa nije napravljeno.")
    else:
        print(f"Gotovo. Napravljeno {ukupno} potvrdnica.")


if __name__ == "__main__":
    main()
